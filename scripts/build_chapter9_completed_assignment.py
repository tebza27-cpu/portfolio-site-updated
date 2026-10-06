from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "static" / "Starting and Growing My Business - Chapter 9 Assignment - Growing My Business.docx"
OUTPUT = ROOT / "static" / "Starting and Growing My Business - Chapter 9 Assignment - Growing My Business - Completed.docx"


def fill_cell(cell, text):
    cell.paragraphs[0].text = text
    for paragraph in cell.paragraphs[1:]:
        paragraph.text = ""


document = Document(SOURCE)
tables = document.tables

tables[0].cell(1, 0).text = "Sifiso Mokhele"
tables[0].cell(1, 1).text = "[Enter date submitted]"
tables[0].cell(1, 2).text = "[Enter action partner]"

fill_cell(
    tables[1].cell(1, 0),
    "At 13, Elder Sitati wanted to attend a school he admired because he saw education as a key to a better life. He rode his father's bicycle for a half-day journey to ask the principal for a place. Although he felt nervous, he asked directly and was later selected as the only student from his primary school to attend.",
)

for row, answer in {
    1: "In five years, I want to run a sustainable small IT services business serving small businesses and professionals with dependable support, cybersecurity, and cloud services.",
    3: "I want to provide useful service, keep learning, and build greater financial stability for myself and my family.",
    4: "I need stronger skills in pricing, bookkeeping, sales, customer agreements, cloud administration, cybersecurity, and protecting client data.",
    5: "[Before submitting: discuss this plan with a family member or friend, then record their name and the feedback they actually give.]",
}.items():
    fill_cell(tables[2].cell(row, 1), answer)

goals = [
    (
        "Interview five small-business prospects about IT support, backups, and security by 31 October 2026.",
        "To confirm which problems customers will pay to solve.",
        "Contact five owners; ask the same questions; record needs, current solutions, and willingness to pay.",
        "Complete by 31 October 2026; review notes each Friday.",
        "[Choose a family member, mentor, or action partner.]",
    ),
    (
        "Prepare three clear starter service packages by 15 November 2026.",
        "To make the offer and price easy to understand.",
        "List scope and costs; price a support package, a backup check, and a security baseline; ask three prospects for feedback.",
        "Draft by 8 November; revise by 15 November 2026.",
        "[Choose a family member, mentor, or action partner.]",
    ),
    (
        "Begin a separate business savings habit; proposed target R300 each week from 9 October 2026.",
        "To build a reserve for tools, training, and unexpected expenses.",
        "Review the weekly budget; transfer an affordable amount after essential expenses; record each deposit.",
        "Start 9 October 2026; review at month-end. Confirm the amount fits my budget.",
        "[Choose a family member, mentor, or action partner.]",
    ),
]
for row, goal in zip((2, 3, 4), goals):
    for column, answer in enumerate(goal):
        fill_cell(tables[3].cell(row, column), answer)

mentor_answers = {
    1: "An experienced small-business owner; an IT or cybersecurity professional; or a business mentor from a local enterprise network.",
    2: "[Choose and name a person who is willing and has my well-being at heart.]",
    3: "I plan to ask by 9 October 2026.",
    4: "Proposed first meeting: 16 October 2026, subject to the mentor's availability.",
    5: "Monthly, with a short progress update between meetings if needed.",
}
for row, answer in mentor_answers.items():
    fill_cell(tables[4].cell(row, 1), answer)

ideas = [
    "Monthly remote IT support",
    "On-site troubleshooting",
    "Small-business Wi-Fi health checks",
    "Managed backup monitoring",
    "Backup restore tests",
    "Cybersecurity baseline reviews",
    "MFA and password setup",
    "Microsoft 365 or Google Workspace administration",
    "Cloud migration support",
    "Computer setup and maintenance",
    "Staff phishing-awareness sessions",
    "Patch and update checks",
]
for index, idea in enumerate(ideas):
    row, column = divmod(index, 2)
    cell = tables[6].cell(row, column)
    prefix = cell.paragraphs[0].text.strip()
    cell.paragraphs[0].text = f"{prefix} {idea}"

fill_cell(
    tables[7].cell(1, 0),
    "1. Monthly remote support package: recurring help for existing small-business customers. 2. Backup monitoring and restore checks: protects an essential business need. 3. Cybersecurity baseline review: a focused, affordable first step into security support. Each extends my current IT services.",
)
fill_cell(
    tables[8].cell(1, 0),
    "Reach more small businesses in my current service area through referrals, LinkedIn, and local business networks.",
)
fill_cell(
    tables[9].cell(1, 0),
    "Research common IT problems, competitor offers and prices, travel costs, and the time each service requires. Interview prospective customers and test a small pilot first. Use a checklist, agreed response times, clear scope, and follow-up notes so existing customers continue receiving reliable service.",
)
fill_cell(
    tables[10].cell(1, 0),
    "Competitors may offer a recognized brand, quick response, monthly contracts, broad on-site coverage, or established customer reviews. I should verify which advantages matter most to local customers.",
)
fill_cell(
    tables[10].cell(1, 1),
    "Offer clear packages and prices, dependable appointment updates, documented checklists, a follow-up after each job, and practical security and backup guidance. Start with work I can deliver well and do not promise 24/7 cover until I can support it.",
)
fill_cell(
    tables[11].cell(1, 0),
    "[Complete after a real conversation: explain the proposed monthly IT health check, record the customer's actual comments, and note whether they expressed interest. Do not submit an assumed response as fact.]",
)
fill_cell(
    tables[12].cell(1, 0),
    "Growing demand for cloud tools, backups, and basic cybersecurity could support a shift toward recurring support. Limited time, cash flow, or too many one-off requests could be obstacles. I would narrow the offer to the services customers request most and expand only when capacity and quality are stable.",
)
fill_cell(
    tables[13].cell(1, 0),
    "Proposed routine: review a simple weekly budget and move an affordable amount into separate business savings. A starting target of R300 per week is a draft estimate; confirm it against my actual income and expenses before committing.",
)

for paragraph in document.paragraphs:
    if "Separate savings for mission, education, retirement or business" in paragraph.text:
        paragraph.text = paragraph.text.replace("\u2610", "\u25c9")
        paragraph.add_run(" (start this week; proposed target R300/week)")
        break

costs = [
    (
        "Unnecessary travel",
        "Use remote support when suitable and group on-site visits by area. Proposed saving: about R400/month; check against actual travel spending.",
    ),
    (
        "Unused or overlapping subscriptions",
        "Review software and cloud plans monthly, cancel duplicates, and compare suitable lower-cost options. Proposed saving: about R250/month; confirm from invoices.",
    ),
]
for row, answers in zip((1, 2), costs):
    for column, answer in enumerate(answers):
        fill_cell(tables[15].cell(row, column), answer)

plan = {
    1: "First, expand services for current small-business customers with a small monthly IT support offer that can include backup and basic security checks. This builds on my core skills and is easier to test than entering a new market.",
    2: "When customer interviews show demand, the pilot can be delivered reliably, and the price covers the time and costs. A useful first checkpoint is three paying pilot customers and a positive margin for three consecutive months.",
    3: "Use my existing computer, internet, and low-cost remote tools; set aside about four hours a week for the pilot. Keep a provisional R2,000 setup limit only if cash flow allows. Seek help for specialist work rather than promising services beyond my capacity.",
    4: "I could underestimate the time, have a quiet sales month, or face a security or data-loss incident. Limit pilot scope, use written approval and secure practices, keep verified backups, track cash weekly, and pause expansion if service quality or cash flow drops.",
    5: "From 9 October 2026, reduce avoidable travel with remote-first support (proposed R400/month) and remove unused subscriptions (proposed R250/month). Check both estimates against actual spending before recording savings.",
}
for row, answer in plan.items():
    fill_cell(tables[16].cell(row, 1), answer)

fill_cell(
    tables[17].cell(1, 0),
    "I feel impressed to grow carefully by listening to customers, improving the services they already need, and managing money honestly before taking on more work.",
)
fill_cell(
    tables[18].cell(1, 0),
    "This week I will contact three potential customers, ask about their IT and backup needs, review two recurring costs, and choose a realistic weekly savings amount.",
)

for paragraph in document.paragraphs:
    if paragraph.text.startswith("How to complete this assignment:"):
        note = paragraph.add_run(
            "\nDRAFT FOR REVIEW: Confirm bracketed personal details, actual conversations, savings habits and amounts, and signatures before submitting. Proposed amounts are estimates."
        )
        note.italic = True
        break

document.save(OUTPUT)
print(OUTPUT)